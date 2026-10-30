import openpyxl

wb = openpyxl.load_workbook('RMS_Authentication_Integration_Test_Completed.xlsx')
ws = wb.active

updates = {
    'TC-INT-051': 'System terminates session, logs out user, and redirects to URL "/login?logout" with message: "Bạn đã đăng xuất."',
    'TC-INT-052': 'System redirects to URL "/register" and shows error: "Không thể hoàn tất yêu cầu. Vui lòng kiểm tra thông tin và thử lại."',
    'TC-INT-053': 'System rejects request, displays countdown timer and shows error: "Không thể hoàn tất yêu cầu. Vui lòng kiểm tra thông tin và thử lại."',
    'TC-INT-054': 'System rejects request, redirects to URL "/register" and shows error: "Không thể hoàn tất yêu cầu. Vui lòng kiểm tra thông tin và thử lại."',
    'TC-INT-055': 'System rejects request, redirects to URL "/forgot-password" and shows error: "Không thể hoàn tất yêu cầu. Vui lòng kiểm tra thông tin và thử lại."',
    'TC-INT-056': 'System sends new 6-digit OTP to email, resets countdown timer and redirects to URL "/reset-password/otp"',
    'TC-INT-057': 'System rejects request, automatically redirects to URL "/forgot-password" and does not change password again.',
    'TC-INT-058': 'System rejects request due to invalid account state, cancels session and redirects to URL "/forgot-password"'
}

for row in range(11, ws.max_row + 1):
    tc_id = ws.cell(row=row, column=1).value
    if tc_id in updates:
        ws.cell(row=row, column=4, value=updates[tc_id])

wb.save('RMS_Authentication_Integration_Test_Completed.xlsx')
wb.save('docs/members/linhdn/RMS_Authentication_Integration_Test_Completed.xlsx')
print('Excel files updated with simplified expected results successfully!')
